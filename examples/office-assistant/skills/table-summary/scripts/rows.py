"""有界读取CSV/XLSX原始记录，拒绝公式、宏、外链及含糊列。"""

from __future__ import annotations

import csv
import re
import sys
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'office/scripts'))
from documents import check_archive, input_path  # noqa: E402

MONEY = re.compile(r'^[+-]?\d{1,12}(?:\.\d{1,6})?$')


def raw_rows(path: Path):
    """返回工作表、行号、常量值；所有工作表和数据行都有上限。"""
    if path.suffix == '.csv':
        csv.field_size_limit(65536)
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.reader(handle)
            for index, values in enumerate(reader, start=1):
                if index > 10001 or len(values) > 256:
                    raise ValueError(f'{path.name}:CSV行或列数超限')
                yield 'CSV', reader.line_num, values
        return
    if path.suffix != '.xlsx':
        raise ValueError('表格仅支持UTF-8 CSV和无公式XLSX')
    import openpyxl

    check_archive(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if not 0 < len(workbook.worksheets) <= 10:
            raise ValueError('XLSX工作表数量须为1至10')
        for sheet in workbook:
            if sheet.max_column and sheet.max_column > 256:
                raise ValueError(f'{path.name}/{sheet.title}:列数超限')
            for index, cells in enumerate(sheet.iter_rows(), start=1):
                if index > 10001 or len(cells) > 256:
                    raise ValueError(f'{path.name}/{sheet.title}:行或列数超限')
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError(f'{path.name}/{sheet.title}:第{index}行含公式，请提供值快照')
                yield sheet.title, index, [cell.value for cell in cells]
    finally:
        workbook.close()


def records(names: list[str], group_column: str, amount_column: str, currency_column: str):
    """定位精确列名，空行忽略；任意无效金额列出来源，不静默补零。"""
    headers = {}
    count = 0
    errors = []
    for name in names:
        path = input_path(name)
        for sheet, row, values in raw_rows(path):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            key = name, sheet
            if key not in headers:
                header = [str(value).strip() if value is not None else '' for value in values]
                wanted = [group_column, amount_column] + ([currency_column] if currency_column else [])
                if any(not column or header.count(column) != 1 for column in wanted):
                    raise ValueError(f'{name}/{sheet}:所需列缺失或重名')
                headers[key] = [header.index(column) for column in wanted]
                continue
            indexes = headers[key]
            try:
                group = str(values[indexes[0]] or '').strip()
                raw = str(values[indexes[1]] if values[indexes[1]] is not None else '').strip()
                currency = str(values[indexes[2]] or '').strip() if currency_column else '未注明'
                if not group or len(group) > 255 or not currency or len(currency) > 64 or not MONEY.fullmatch(raw):
                    raise ValueError('分组、币种为空或金额无效（最多12位整数和6位小数）')
                amount = Decimal(raw)
            except (IndexError, ValueError) as exc:
                errors.append({'file':name,'sheet':sheet,'row':row,'error':str(exc)})
                if len(errors) > 100:
                    raise ValueError('无效行超过100条，请先修正输入')
                continue
            count += 1
            if count > 10000:
                raise ValueError('单任务原始记录总数超过10000行')
            yield {'file':name,'sheet':sheet,'row':row,'group':group,'currency':currency,'amount':amount}
    if errors:
        import json

        raise ValueError('无效原始行：' + json.dumps(errors, ensure_ascii=False))
    if count == 0:
        raise ValueError('表格没有可汇总的有效原始记录')
